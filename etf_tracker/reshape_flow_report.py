# -*- coding: utf-8 -*-
"""将现有四Sheet分类流动报告重构为月度矩阵表和降序周度图。"""
import os, json
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "ETF净流动_按用户四Sheet分类_2026.xlsx")
OUT = os.path.join(BASE, "ETF净流动_按用户四Sheet分类_2026_矩阵版.xlsx")
HTML = os.path.join(BASE, "ETF净流动_按用户四Sheet分类_2026_矩阵版.html")
SHEETS = ["宽基ETF", "行业ETF", "风格ETF", "港股ETF"]
PREFERRED = {
    "宽基ETF": ["上证50","沪深300","中证A500","科创板","中证1000","中证500","中证2000"]
}
monthly = pd.read_excel(SRC, sheet_name="全部月度净流动")
weekly = pd.read_excel(SRC, sheet_name="全部最新一周")
week_periods = weekly["统计区间"].dropna().astype(str).unique().tolist()
week_period = week_periods[0] if week_periods else "暂无有效区间"
mapping = pd.read_excel(SRC, sheet_name="全部分类映射", dtype={"基金代码":str})
detail = pd.read_excel(SRC, sheet_name="全部ETF明细", dtype={"基金代码":str})
daily_week = pd.read_excel(SRC, sheet_name="最新一周逐日净流动")

matrices = {}
daily_matrices = {}
weekly_sorted = {}
for sheet in SHEETS:
    sm = monthly[monthly["主分类"] == sheet].copy()
    cats = mapping[mapping["主分类"] == sheet]["分类"].drop_duplicates().tolist()
    if sheet in PREFERRED:
        cats = [c for c in PREFERRED[sheet] if c in cats] + [c for c in cats if c not in PREFERRED[sheet]]
    pivot = sm.pivot_table(index="月份", columns="分类", values="净流动(亿元)", aggfunc="sum", fill_value=0)
    pivot = pivot.reindex(columns=cats, fill_value=0)
    pivot.index = [f"{str(x)[:4]}年{int(str(x)[5:7])}月" for x in pivot.index]
    pivot.loc["2026年合计"] = pivot.sum(axis=0)
    pivot.index.name = "月份"
    matrices[sheet] = pivot
    dw = daily_week[daily_week["主分类"] == sheet].pivot_table(index="日期", columns="分类", values="净流动(亿元)", aggfunc="sum", fill_value=0)
    dw = dw.reindex(columns=cats, fill_value=0)
    dw.index = [str(x)[:10] for x in dw.index]
    dw.index.name = "交易日"
    daily_matrices[sheet] = dw
    weekly_sorted[sheet] = weekly[weekly["主分类"] == sheet].sort_values("净流动(亿元)", ascending=False).copy()

with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
    for sheet in SHEETS:
        matrices[sheet].to_excel(writer, sheet_name=f"{sheet[:2]}月度矩阵")
        daily_matrices[sheet].to_excel(writer, sheet_name=f"{sheet[:2]}周度逐日矩阵")
        weekly_sorted[sheet][["统计区间","分类","净流动(亿元)","ETF数量"]].to_excel(writer,index=False,sheet_name=f"{sheet[:2]}周度降序")
    monthly.to_excel(writer,index=False,sheet_name="月度长表")
    weekly.to_excel(writer,index=False,sheet_name="周度长表")
    mapping.to_excel(writer,index=False,sheet_name="分类映射")
    detail.to_excel(writer,index=False,sheet_name="ETF明细")
    wb=writer.book
    red_fill=PatternFill("solid",fgColor="FCE8E6"); green_fill=PatternFill("solid",fgColor="E6F4EA")
    for ws in wb.worksheets:
        ws.freeze_panes="B2" if "矩阵" in ws.title else "A2"
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill("solid",fgColor="D9E7F5"); cell.font=Font(bold=True); cell.alignment=Alignment(horizontal="center")
        for col in range(1,ws.max_column+1):
            max_len=max(len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,100)+1))
            ws.column_dimensions[get_column_letter(col)].width=min(max(max_len+2,12),26)
        if "矩阵" in ws.title:
            for cell in ws[ws.max_row]: cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="FFF2CC")
            rng=f"B2:{get_column_letter(ws.max_column)}{ws.max_row}"
            ws.conditional_formatting.add(rng,CellIsRule(operator="greaterThan",formula=["0"],fill=red_fill))
            ws.conditional_formatting.add(rng,CellIsRule(operator="lessThan",formula=["0"],fill=green_fill))

nav="".join(f'<button onclick="showTab(\'{s}\')">{s}</button>' for s in SHEETS)
sections=[]
for idx,sheet in enumerate(SHEETS):
    matrix=matrices[sheet]
    cols=list(matrix.columns)
    header="".join(f"<th>{c}</th>" for c in cols)
    trs=[]
    for label,row in matrix.iterrows():
        cells="".join(f'<td class="{"pos" if v>0 else "neg" if v<0 else ""}">{float(v):+.2f}</td>' for v in row)
        trs.append(f'<tr class="{"total" if label=="2026年合计" else ""}"><th>{label}</th>{cells}</tr>')
    w=weekly_sorted[sheet]
    cats=w["分类"].astype(str).tolist(); vals=[round(float(x),2) for x in w["净流动(亿元)"]]
    dw = daily_matrices[sheet]
    dheader="".join(f"<th>{c}</th>" for c in dw.columns)
    dtrs=[]
    for label,row in dw.iterrows():
        cells="".join(f'<td class="{"pos" if v>0 else "neg" if v<0 else ""}">{float(v):+.2f}</td>' for v in row)
        dtrs.append(f'<tr><th>{label}</th>{cells}</tr>')
    sections.append(f'''<section id="{sheet}" class="tab"><div class="panel"><h2>{sheet}：2026年月度净流动矩阵（亿元）</h2><div class="scroll"><table class="matrix"><thead><tr><th>月份</th>{header}</tr></thead><tbody>{''.join(trs)}</tbody></table></div></div><div class="panel"><h2>{sheet}：最新一周净流动（{week_period}，从大到小）</h2><div class="chart"><canvas id="w{idx}"></canvas></div></div><div class="panel"><h2>{sheet}：最新一周逐日净流动矩阵（亿元）</h2><div class="scroll"><table class="matrix"><thead><tr><th>交易日</th>{dheader}</tr></thead><tbody>{''.join(dtrs)}</tbody></table></div></div><script type="application/json" id="d{idx}">{json.dumps({'cats':cats,'vals':vals},ensure_ascii=False)}</script></section>''')
html=f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>ETF净流动矩阵报告</title><script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.0/dist/chart.umd.min.js"></script><style>body{{font-family:Arial,"Microsoft YaHei";background:#f5f7fb;color:#222;margin:0}}.wrap{{max-width:1450px;margin:auto;padding:24px}}.hero,.panel{{background:#fff;border-radius:12px;padding:22px;margin-bottom:18px;box-shadow:0 2px 10px #dfe5ef}}button{{padding:9px 18px;margin:4px;border:0;border-radius:7px;background:#e7eef8;cursor:pointer}}button.active{{background:#2166ac;color:white}}.tab{{display:none}}.tab.active{{display:block}}.chart{{height:520px}}.scroll{{overflow-x:auto}}table{{width:100%;border-collapse:collapse;font-size:13px}}th,td{{padding:9px;border:1px solid #e1e7ef;text-align:right;white-space:nowrap}}th{{background:#eef3fa;text-align:center}}.matrix tbody th{{text-align:left}}.pos{{color:#d32f2f;font-weight:700;background:#fff4f2}}.neg{{color:#238636;font-weight:700;background:#f1fbf4}}tr.total th,tr.total td{{font-weight:800;background:#fff4cc}}.note{{font-size:13px;color:#65748b}}</style></head><body><div class="wrap"><div class="hero"><h1>ETF净流动分类矩阵</h1><p class="note">数据最后更新：{pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}。最新有效交易周：{week_period}。月度数据采用矩阵表，周度图按净流动从大到小排列。正数为净流入（红色），负数为净流出（绿色）。</p><div>{nav}</div></div>{''.join(sections)}<div class="panel note">分类采用用户ETF分类.xlsx。数据源为交易所每日原始份额及恒生聚源单位净值NAV；净流动=(当日原始份额-上一交易日原始份额)×当日NAV；周度严格按周一至周五实际交易日。每周六更新。</div></div><script>function showTab(id){{document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));document.getElementById(id).classList.add('active');document.querySelectorAll('button').forEach(x=>x.classList.toggle('active',x.textContent===id));}}document.querySelectorAll('.tab').forEach((sec,i)=>{{const d=JSON.parse(document.getElementById('d'+i).textContent);new Chart(document.getElementById('w'+i),{{type:'bar',data:{{labels:d.cats,datasets:[{{data:d.vals,backgroundColor:d.vals.map(v=>v>=0?'rgba(211,47,47,.75)':'rgba(35,134,54,.75)')}}]}},options:{{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{{legend:{{display:false}},tooltip:{{callbacks:{{label:c=>(c.raw>=0?'净流入 ':'净流出 ')+Math.abs(c.raw).toFixed(2)+' 亿元'}}}}}},scales:{{x:{{title:{{display:true,text:'净流动（亿元）'}}}}}}}}}});}});showTab('宽基ETF');</script></body></html>'''
with open(HTML,"w",encoding="utf-8") as f:f.write(html)
print(OUT);print(HTML)
for s in SHEETS:
    print(s,matrices[s].shape,'weekly_desc',weekly_sorted[s]["净流动(亿元)"].is_monotonic_decreasing)
